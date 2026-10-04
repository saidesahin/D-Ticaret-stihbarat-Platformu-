"""
ORTAK ALTYAPI - excel_export.py
Dokumanda hemen her modulun ciktisi "Excel'e aktar" oldugu icin
bunu ortak bir fonksiyon olarak core'a koyuyoruz. Modul 1, 2, 3...
hicbiri kendi Excel kodunu yazmaz, hepsi bunu cagirir.
"""
import io
from typing import Sequence

import openpyxl
from openpyxl.utils import get_column_letter
from fastapi.responses import StreamingResponse


def rows_to_excel_response(rows: Sequence[dict], sheet_name: str, filename: str) -> StreamingResponse:
    """
    rows: [{"Sutun1": deger, "Sutun2": deger, ...}, ...]
    Sozluklerin key'leri otomatik olarak baslik satiri olur.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet adi 31 karakter siniri

    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])

        for i, header in enumerate(headers, start=1):
            col_letter = get_column_letter(i)
            max_len = max([len(str(header))] + [len(str(r.get(header, ""))) for r in rows])
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)
    else:
        ws.append(["Kayit bulunamadi"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )
